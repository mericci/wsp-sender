import openpyxl

#doc = openpyxl.load_workbook('BD DMT 2018-1.xlsx') #Acá va el archivo que voy a leer
doc = openpyxl.load_workbook('Invitacion MdV.xlsx') #OJO ACA

i = 0
ws = doc.get_sheet_by_name('Hoja2')  #esto era hoja2
for elem in ws: #recorro hoja 2

    print(elem[3].value)

    i +=1
    if elem[0].value == "Nombre":

        continue

    numero = elem[3].value
    print(numero)
    numero = str(numero)
    numero = numero[-8:]
    numero = "+569" + numero
    ws.cell(row=i, column=4).value = numero

doc.save("WSPS.xlsx")